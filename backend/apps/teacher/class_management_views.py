"""Teacher teaching-class creation and student-import APIs."""

from __future__ import annotations

import uuid

from openpyxl import load_workbook
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response

from api.permissions import IsTeacher, session_user
from apps.context.services import CurrentClassService
from repositories.learning_repository import LearningRepository
from repositories.mysql_user_repository import MySQLUserRepository
from repositories.teacher_class_repository import TeacherClassRepository


_PENDING_IMPORTS: dict[str, dict] = {}
_GENDER_CODES = {"男": 1, "女": 2, "其他": 0, "": None}


def _error(message: str, code: str, code_status: int):
    return Response({"message": message, "code": code}, status=code_status)


@api_view(["GET"])
@permission_classes([IsTeacher])
def classes(request):
    with TeacherClassRepository() as repository:
        items = repository.list_owned_classes(session_user(request)["username"])
    return Response({"items": items, "meta": {"total": len(items)}})


@api_view(["POST"])
@permission_classes([IsTeacher])
def create_class(request):
    title = str(request.data.get("title", "")).strip()
    teaching_class = str(request.data.get("teaching_class", "")).strip()
    academic_year = str(request.data.get("academic_year", "")).strip()
    if not title or not teaching_class:
        return _error("课程名称和教学班名称不能为空。", "CLASS_INVALID", status.HTTP_400_BAD_REQUEST)
    if len(title) > 255 or len(teaching_class) > 128 or len(academic_year) > 32:
        return _error("教学班信息长度超过限制。", "CLASS_VALUE_TOO_LONG", status.HTTP_400_BAD_REQUEST)
    username = session_user(request)["username"]
    try:
        with TeacherClassRepository() as repository:
            existing = repository.find_existing(title, teaching_class, academic_year)
            if existing:
                return Response(
                    {"message": "教学班已存在，请直接使用已有教学班。", "code": "CLASS_EXISTS", "class": existing},
                    status=status.HTTP_409_CONFLICT,
                )
            item = repository.create_class_for_teacher(username, title, teaching_class, academic_year)
        CurrentClassService().select(request, str(item["id"]))
        with LearningRepository() as audit:
            audit.write_audit(session_user(request), "class.create", "classes", str(item.get("id", "")), {"teacher_username": username, **item})
        return Response({"class": item}, status=status.HTTP_201_CREATED)
    except Exception as exc:
        if "Duplicate" in str(exc) or "1062" in str(exc):
            return _error("教学班已存在，请直接使用已有教学班。", "CLASS_EXISTS", status.HTTP_409_CONFLICT)
        return _error("教学班创建失败。", "CLASS_CREATE_FAILED", status.HTTP_400_BAD_REQUEST)


@api_view(["POST"])
@permission_classes([IsTeacher])
def import_students_preview(request):
    uploaded = request.FILES.get("file")
    class_id = str(request.data.get("class_id", "")).strip()
    if uploaded is None:
        return _error("请选择 Excel 文件。", "IMPORT_FILE_REQUIRED", status.HTTP_400_BAD_REQUEST)
    if uploaded.size > 5 * 1024 * 1024:
        return _error("Excel 文件不能超过 5MB。", "IMPORT_FILE_TOO_LARGE", status.HTTP_400_BAD_REQUEST)
    if not class_id:
        return _error("请先选择教学班。", "CLASS_REQUIRED", status.HTTP_400_BAD_REQUEST)
    username = session_user(request)["username"]
    try:
        with TeacherClassRepository() as repository:
            teaching_class_item = repository.get_owned_class(username, class_id)
        if not teaching_class_item:
            return _error("只能向本人管理的启用教学班导入学生。", "CLASS_FORBIDDEN", status.HTTP_403_FORBIDDEN)
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
        sheet = workbook["账号导入"] if "账号导入" in workbook.sheetnames else workbook.active
        values = list(sheet.iter_rows(values_only=True))
        header = list(values[3]) if len(values) > 3 else []
        required = ["学号", "姓名", "班级名称", "性别"]
        if [str(value or "").strip() for value in header[:4]] != required:
            return _error("模板表头不正确，请使用系统提供的模板。", "IMPORT_HEADER_INVALID", status.HTTP_400_BAD_REQUEST)
        rows, errors, seen = [], [], set()
        with MySQLUserRepository() as account_repository:
            for row_number, values_row in enumerate(values[4:], start=5):
                values_row = list(values_row) + [None] * 4
                if not any(str(value or "").strip() for value in values_row[:4]):
                    continue
                student_username = str(values_row[0] or "").strip()
                name = str(values_row[1] or "").strip()
                gender = str(values_row[3] or "").strip()
                item = {
                    "row": row_number,
                    "username": student_username,
                    "name": name,
                    "gender_label": gender,
                    "administrative_class": str(values_row[2] or "").strip(),
                    "teaching_class": teaching_class_item["teaching_class"],
                }
                row_errors = []
                if not student_username:
                    row_errors.append("用户名（学号）不能为空")
                if not name:
                    row_errors.append("姓名不能为空")
                if gender not in _GENDER_CODES:
                    row_errors.append("性别只能填写男、女或其他")
                if student_username in seen:
                    row_errors.append("文件内用户名重复")
                elif student_username and account_repository.username_exists(student_username):
                    row_errors.append("用户名已存在，请联系管理员添加到本教学班")
                seen.add(student_username)
                if row_errors:
                    item["errors"] = row_errors
                    errors.append(item)
                else:
                    rows.append({**item, "gender": _GENDER_CODES[gender]})
        token = uuid.uuid4().hex
        _PENDING_IMPORTS[token] = {"teacher_username": username, "class_id": class_id, "rows": rows}
        return Response({
            "token": token,
            "class": teaching_class_item,
            "rows": rows,
            "errors": errors,
            "can_import": bool(rows) and not errors,
            "meta": {"valid": len(rows), "invalid": len(errors)},
        })
    except Exception:
        return _error("Excel 解析失败，请确认文件格式正确。", "IMPORT_PARSE_FAILED", status.HTTP_400_BAD_REQUEST)


@api_view(["POST"])
@permission_classes([IsTeacher])
def import_students_confirm(request):
    token = str(request.data.get("token", ""))
    pending = _PENDING_IMPORTS.get(token)
    if not pending:
        return _error("导入预览已失效，请重新上传文件。", "IMPORT_PREVIEW_EXPIRED", status.HTTP_400_BAD_REQUEST)
    username = session_user(request)["username"]
    if pending["teacher_username"] != username:
        return _error("无权确认这批导入数据。", "IMPORT_FORBIDDEN", status.HTTP_403_FORBIDDEN)
    try:
        with TeacherClassRepository() as repository:
            created = repository.import_students(username, pending["class_id"], pending["rows"])
        with LearningRepository() as audit:
            audit.write_audit(session_user(request), "class.students.import", "classes", pending["class_id"], {"count": created})
        _PENDING_IMPORTS.pop(token, None)
        return Response({"created": created, "class_id": pending["class_id"]}, status=status.HTTP_201_CREATED)
    except PermissionError:
        return _error("只能向本人管理的启用教学班导入学生。", "CLASS_FORBIDDEN", status.HTTP_403_FORBIDDEN)
    except Exception as exc:
        if "Duplicate" in str(exc) or "1062" in str(exc):
            return _error("用户名已存在，请联系管理员添加到本教学班。", "USERNAME_EXISTS", status.HTTP_409_CONFLICT)
        return _error("批量导入失败，数据库未完成本次导入。", "IMPORT_CREATE_FAILED", status.HTTP_400_BAD_REQUEST)


@api_view(["POST"])
@permission_classes([IsTeacher])
def create_student(request):
    class_id = str(request.data.get("class_id", "")).strip()
    username = str(request.data.get("username", "")).strip()
    name = str(request.data.get("name", "")).strip()
    administrative_class = str(request.data.get("administrative_class", request.data.get("study_class", ""))).strip()
    raw_gender = request.data.get("gender", "")
    gender_label = str(request.data.get("gender_label", "")).strip()
    if not class_id:
        return _error("请先选择教学班。", "CLASS_REQUIRED", status.HTTP_400_BAD_REQUEST)
    if not username or not name:
        return _error("学号和姓名不能为空。", "STUDENT_INVALID", status.HTTP_400_BAD_REQUEST)
    if len(username) > 64 or len(name) > 128 or len(administrative_class) > 128:
        return _error("学生信息长度超过限制。", "STUDENT_VALUE_TOO_LONG", status.HTTP_400_BAD_REQUEST)
    if gender_label:
        if gender_label not in _GENDER_CODES:
            return _error("性别只能填写男、女或其他。", "STUDENT_GENDER_INVALID", status.HTTP_400_BAD_REQUEST)
        gender = _GENDER_CODES[gender_label]
    else:
        try:
            gender = None if raw_gender in (None, "") else int(raw_gender)
        except (TypeError, ValueError):
            return _error("性别参数无效。", "STUDENT_GENDER_INVALID", status.HTTP_400_BAD_REQUEST)
        if gender not in {None, 0, 1, 2}:
            return _error("性别参数无效。", "STUDENT_GENDER_INVALID", status.HTTP_400_BAD_REQUEST)
    teacher_username = session_user(request)["username"]
    try:
        with TeacherClassRepository() as repository:
            class_item = repository.get_owned_class(teacher_username, class_id)
        if not class_item:
            return _error("只能向本人管理的启用教学班添加学生。", "CLASS_FORBIDDEN", status.HTTP_403_FORBIDDEN)
        with MySQLUserRepository() as account_repository:
            if account_repository.username_exists(username):
                return _error("用户名已存在，请联系管理员添加到本教学班。", "USERNAME_EXISTS", status.HTTP_409_CONFLICT)
        row = {
            "username": username,
            "name": name,
            "gender": gender,
            "administrative_class": administrative_class,
            "teaching_class": class_item["teaching_class"],
        }
        with TeacherClassRepository() as repository:
            repository.create_student(teacher_username, class_id, row)
        with LearningRepository() as audit:
            audit.write_audit(session_user(request), "class.student.create", "classes", class_id, {"username": username})
        return Response({"created": 1, "student": {"username": username, "name": name, "class_id": class_id}}, status=status.HTTP_201_CREATED)
    except PermissionError:
        return _error("只能向本人管理的启用教学班添加学生。", "CLASS_FORBIDDEN", status.HTTP_403_FORBIDDEN)
    except Exception as exc:
        if "Duplicate" in str(exc) or "1062" in str(exc):
            return _error("用户名已存在，请联系管理员添加到本教学班。", "USERNAME_EXISTS", status.HTTP_409_CONFLICT)
        return _error("学生账号创建失败。", "STUDENT_CREATE_FAILED", status.HTTP_400_BAD_REQUEST)
