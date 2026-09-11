"""Administrator account, permission and audit APIs (F24/F25)."""

import uuid

from openpyxl import load_workbook

from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response

from api.permissions import IsAdmin, session_user
from repositories.admin_class_repository import AdminClassRepository
from repositories.learning_repository import LearningRepository
from repositories.mysql_user_repository import MySQLUserRepository


_PENDING_IMPORTS: dict[str, list[dict]] = {}


def _error(message: str, code: str, code_status: int):
    return Response({"message": message, "code": code}, status=code_status)


def _boolean(value):
    if isinstance(value, bool):
        return value
    if value in (1, "1", "true", "True"):
        return True
    if value in (0, "0", "false", "False"):
        return False
    raise ValueError("布尔值无效")


def _accounts(repository: MySQLUserRepository, role: str):
    if role == "admin":
        return repository.list_admins()
    if role == "teacher":
        return repository.list_teachers()
    if role == "student":
        return repository.list_students()
    return repository.list_admins() + repository.list_teachers() + repository.list_students()


@api_view(["GET"])
@permission_classes([IsAdmin])
def accounts(request):
    role = request.query_params.get("role", "all")
    if role not in {"all", "admin", "teacher", "student"}:
        return _error("角色参数无效。", "ROLE_INVALID", status.HTTP_400_BAD_REQUEST)
    with MySQLUserRepository() as repository:
        items = _accounts(repository, role)
    available_administrative_classes = sorted({str(item.get("administrative_class", "")).strip() for item in items if str(item.get("administrative_class", "")).strip()})
    available_teaching_classes = sorted({str(item.get("teaching_class", "")).strip() for item in items if str(item.get("teaching_class", "")).strip()})

    keyword = request.query_params.get("q", "").strip().lower()
    administrative_class = request.query_params.get("administrative_class", request.query_params.get("study_class", "")).strip()
    teaching_class = request.query_params.get("teaching_class", "").strip()
    account_status = request.query_params.get("status", "all")
    if keyword:
        items = [item for item in items if keyword in str(item.get("username", "")).lower() or keyword in str(item.get("name", "")).lower()]
    if administrative_class:
        items = [item for item in items if item.get("administrative_class", "") == administrative_class]
    if teaching_class:
        items = [item for item in items if item.get("teaching_class", "") == teaching_class]
    if account_status in {"active", "inactive"}:
        expected = "启用" if account_status == "active" else "停用"
        items = [item for item in items if item.get("status") == expected]

    total = len(items)
    role_counts = {key: sum(item.get("role") == key for item in items) for key in ("student", "teacher", "admin")}
    status_counts = {key: sum(item.get("status") == value for item in items) for key, value in (("active", "启用"), ("inactive", "停用"))}
    try:
        page = max(1, int(request.query_params.get("page", 1)))
        page_size = min(100, max(1, int(request.query_params.get("page_size", 20))))
    except (TypeError, ValueError):
        return _error("分页参数无效。", "PAGINATION_INVALID", status.HTTP_400_BAD_REQUEST)
    start = (page - 1) * page_size
    return Response({
        "items": items[start:start + page_size],
        "meta": {
            "role": role,
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": (total + page_size - 1) // page_size,
            "administrative_classes": available_administrative_classes,
            "teaching_classes": available_teaching_classes,
            "role_counts": role_counts,
            "status_counts": status_counts,
        },
    })


@api_view(["POST"])
@permission_classes([IsAdmin])
def create_account(request):
    role = str(request.data.get("role", "student"))
    username = str(request.data.get("username", "")).strip()
    if role not in {"admin", "teacher", "student"} or not username:
        return _error("账号角色和用户名不能为空。", "ACCOUNT_INVALID", status.HTTP_400_BAD_REQUEST)
    class_ids = request.data.get("class_ids", [])
    if class_ids is None:
        class_ids = []
    if not isinstance(class_ids, list):
        class_ids = [class_ids]
    if role not in {"teacher", "student"}:
        class_ids = []
    try:
        with AdminClassRepository() as class_repository:
            valid_class_ids = class_repository.valid_class_ids(class_ids)
        invalid_class_ids = [str(value) for value in class_ids if str(value).strip() and str(value) not in valid_class_ids]
        if invalid_class_ids:
            return Response(
                {"message": "存在无效或已停用的教学班，请重新选择。", "code": "CLASS_INVALID", "class_ids": invalid_class_ids},
                status=status.HTTP_400_BAD_REQUEST,
            )
        with MySQLUserRepository() as repository:
            creator = {"admin": repository.create_admin, "teacher": repository.create_teacher, "student": repository.create_student}[role]
            item = creator({**request.data, "username": username})
        if valid_class_ids:
            with AdminClassRepository() as class_repository:
                class_repository.replace_account_classes(role, username, valid_class_ids)
        with LearningRepository() as audit:
            audit.write_audit(session_user(request), "account.create", role, username, {"username": username, "class_ids": valid_class_ids})
    except Exception as exc:
        if "Duplicate" in str(exc) or "1062" in str(exc):
            return _error("用户名已存在。", "USERNAME_EXISTS", status.HTTP_409_CONFLICT)
        return _error("账号创建失败。", "ACCOUNT_CREATE_FAILED", status.HTTP_400_BAD_REQUEST)
    return Response({"account": item}, status=status.HTTP_201_CREATED)


@api_view(["PATCH"])
@permission_classes([IsAdmin])
def update_account(request, role: str, username: str):
    if role not in {"admin", "teacher", "student"}:
        return _error("角色参数无效。", "ROLE_INVALID", status.HTTP_400_BAD_REQUEST)

    profile_data = {}
    if "name" in request.data:
        name = str(request.data.get("name", "")).strip()
        if not name:
            return _error("姓名不能为空。", "ACCOUNT_NAME_REQUIRED", status.HTTP_400_BAD_REQUEST)
        if len(name) > 100:
            return _error("姓名不能超过 100 个字符。", "ACCOUNT_NAME_TOO_LONG", status.HTTP_400_BAD_REQUEST)
        profile_data["name"] = name
    if role == "student" and "gender" in request.data:
        raw_gender = request.data.get("gender")
        try:
            gender = None if raw_gender in (None, "") else int(raw_gender)
        except (TypeError, ValueError):
            return _error("性别参数无效。", "ACCOUNT_GENDER_INVALID", status.HTTP_400_BAD_REQUEST)
        if gender not in {None, 0, 1, 2}:
            return _error("性别参数无效。", "ACCOUNT_GENDER_INVALID", status.HTTP_400_BAD_REQUEST)
        profile_data["gender"] = gender
    if role == "student" and ("administrative_class" in request.data or "study_class" in request.data):
        administrative_class = str(request.data.get("administrative_class", request.data.get("study_class", ""))).strip()
        if len(administrative_class) > 100:
            return _error("行政班名称不能超过 100 个字符。", "ACCOUNT_CLASS_TOO_LONG", status.HTTP_400_BAD_REQUEST)
        profile_data["study_class"] = administrative_class

    active = None
    if "is_active" in request.data or "active" in request.data:
        try:
            active = _boolean(request.data.get("is_active", request.data.get("active")))
        except ValueError:
            return _error("账号状态参数无效。", "ACCOUNT_STATUS_INVALID", status.HTTP_400_BAD_REQUEST)

    class_ids = None
    valid_class_ids: list[str] = []
    if "class_ids" in request.data and role in {"teacher", "student"}:
        class_ids = request.data.get("class_ids")
        if not isinstance(class_ids, list):
            class_ids = [class_ids] if class_ids is not None else []
        with AdminClassRepository() as class_repository:
            valid_class_ids = class_repository.valid_class_ids(class_ids)
        invalid_class_ids = [str(value) for value in class_ids if str(value).strip() and str(value) not in valid_class_ids]
        if invalid_class_ids:
            return _error("部分教学班无效或已停用，未保存这些关系。", "CLASS_INVALID", status.HTTP_400_BAD_REQUEST)

    updated_fields: list[str] = []
    with MySQLUserRepository() as repository:
        if not repository.account_exists(role, username):
            return _error("账号不存在。", "ACCOUNT_NOT_FOUND", status.HTTP_404_NOT_FOUND)
        if profile_data:
            repository.update_profile(role, username, profile_data)
            updated_fields.extend(profile_data.keys())
        if active is not None:
            repository.set_active(role, username, active)
            updated_fields.append("is_active")
        if request.data.get("password"):
            repository.reset_password(role, username, str(request.data["password"]))
            updated_fields.append("password")
    if class_ids is not None:
        with AdminClassRepository() as class_repository:
            class_repository.replace_account_classes(role, username, valid_class_ids)
        updated_fields.append("class_ids")
    if not updated_fields:
        return _error("没有可更新的账号字段。", "ACCOUNT_UPDATE_INVALID", status.HTTP_400_BAD_REQUEST)
    with LearningRepository() as audit:
        audit.write_audit(session_user(request), "account.update", role, username, {"fields": updated_fields})
    return Response({"updated": True, "role": role, "username": username, "fields": updated_fields})


@api_view(["POST"])
@permission_classes([IsAdmin])
def import_accounts_preview(request):
    uploaded = request.FILES.get("file")
    if uploaded is None:
        return _error("请选择 Excel 文件。", "IMPORT_FILE_REQUIRED", status.HTTP_400_BAD_REQUEST)
    if uploaded.size > 5 * 1024 * 1024:
        return _error("Excel 文件不能超过 5MB。", "IMPORT_FILE_TOO_LARGE", status.HTTP_400_BAD_REQUEST)
    teaching_class = str(request.data.get("teaching_class", "")).strip()
    if not teaching_class:
        return _error("请填写本批次统一使用的教学班。", "TEACHING_CLASS_REQUIRED", status.HTTP_400_BAD_REQUEST)
    try:
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
        sheet = workbook["账号导入"] if "账号导入" in workbook.sheetnames else workbook.active
        values = list(sheet.iter_rows(values_only=True))
        header = list(values[3]) if len(values) > 3 else []
        required = ["学号", "姓名", "班级名称", "性别"]
        if [str(value or "").strip() for value in header[:4]] != required:
            return _error("模板表头不正确，请使用系统提供的模板。", "IMPORT_HEADER_INVALID", status.HTTP_400_BAD_REQUEST)
        rows, errors, seen = [], [], set()
        gender_codes = {"男": 1, "女": 2, "其他": 0, "": None}
        with MySQLUserRepository() as repository:
            for row_number, values_row in enumerate(values[4:], start=5):
                values_row = list(values_row) + [None] * 4
                if not any(str(value or "").strip() for value in values_row[:4]):
                    continue
                username = str(values_row[0] or "").strip()
                name = str(values_row[1] or "").strip()
                gender = str(values_row[3] or "").strip()
                item = {"row": row_number, "username": username, "name": name, "gender_label": gender, "administrative_class": str(values_row[2] or "").strip(), "teaching_class": teaching_class}
                row_errors = []
                if not username: row_errors.append("用户名（学号）不能为空")
                if not name: row_errors.append("姓名不能为空")
                if gender not in gender_codes: row_errors.append("性别只能填写男、女或其他")
                if username in seen: row_errors.append("文件内用户名重复")
                elif username and repository.username_exists(username): row_errors.append("用户名已存在")
                seen.add(username)
                if row_errors:
                    item["errors"] = row_errors; errors.append(item)
                else:
                    rows.append({**item, "gender": gender_codes[gender]})
        token = uuid.uuid4().hex
        _PENDING_IMPORTS[token] = rows
        return Response({"token": token, "rows": rows, "errors": errors, "can_import": bool(rows) and not errors, "meta": {"valid": len(rows), "invalid": len(errors)}})
    except Exception:
        return _error("Excel 解析失败，请确认文件格式正确。", "IMPORT_PARSE_FAILED", status.HTTP_400_BAD_REQUEST)


@api_view(["POST"])
@permission_classes([IsAdmin])
def import_accounts_confirm(request):
    token = str(request.data.get("token", ""))
    rows = _PENDING_IMPORTS.get(token)
    if not rows:
        return _error("导入预览已失效，请重新上传文件。", "IMPORT_PREVIEW_EXPIRED", status.HTTP_400_BAD_REQUEST)
    try:
        with MySQLUserRepository() as repository:
            accounts = repository.create_students_batch(rows)
        with LearningRepository() as audit:
            audit.write_audit(session_user(request), "account.import", "student", "batch", {"count": len(accounts)})
        _PENDING_IMPORTS.pop(token, None)
        return Response({"created": len(accounts), "accounts": accounts}, status=status.HTTP_201_CREATED)
    except Exception:
        return _error("批量导入失败，数据库未完成本次导入。", "IMPORT_CREATE_FAILED", status.HTTP_400_BAD_REQUEST)


@api_view(["GET"])
@permission_classes([IsAdmin])
def audit_logs(request):
    with LearningRepository() as repository:
        items = repository.list_audits(int(request.query_params.get("limit", 100)))
    return Response({"items": items, "meta": {"total": len(items)}})


@api_view(["GET"])
@permission_classes([IsAdmin])
def classes(request):
    keyword = request.query_params.get("q", "").strip()
    class_status = request.query_params.get("status", "all")
    if class_status not in {"all", "active", "inactive"}:
        return _error("班级状态参数无效。", "CLASS_STATUS_INVALID", status.HTTP_400_BAD_REQUEST)
    with AdminClassRepository() as repository:
        items = repository.list_classes(keyword, class_status)
        teachers = repository.account_options("teacher")
        students = repository.account_options("student")
    return Response({"items": items, "meta": {"total": len(items), "teachers": teachers, "students": students}})


@api_view(["POST"])
@permission_classes([IsAdmin])
def create_class(request):
    title = str(request.data.get("title", "")).strip()
    teaching_class = str(request.data.get("teaching_class", "")).strip()
    academic_year = str(request.data.get("academic_year", "")).strip()
    if not title or not teaching_class:
        return _error("课程名称和教学班名称不能为空。", "CLASS_INVALID", status.HTTP_400_BAD_REQUEST)
    try:
        with AdminClassRepository() as repository:
            existing = repository.find_existing(title, teaching_class, academic_year)
            if existing:
                return Response({"message": "教学班已存在，可直接选择该班级。", "code": "CLASS_EXISTS", "class": existing}, status=status.HTTP_409_CONFLICT)
            item = repository.create_class(title, teaching_class, academic_year)
        with LearningRepository() as audit:
            audit.write_audit(session_user(request), "class.create", "classes", str(item.get("id", "")), item)
        return Response({"class": item}, status=status.HTTP_201_CREATED)
    except Exception:
        return _error("教学班创建失败。", "CLASS_CREATE_FAILED", status.HTTP_400_BAD_REQUEST)


@api_view(["PATCH"])
@permission_classes([IsAdmin])
def update_class(request, class_id: str):
    supported = {"title", "teaching_class", "academic_year", "is_active"}
    if not any(field in request.data for field in supported):
        return _error("没有可更新的班级字段。", "CLASS_UPDATE_INVALID", status.HTTP_400_BAD_REQUEST)
    with AdminClassRepository() as repository:
        current = repository.get_class(class_id)
        if not current:
            return _error("教学班不存在。", "CLASS_NOT_FOUND", status.HTTP_404_NOT_FOUND)
        update_data = {}
        for field, limit in (("title", 200), ("teaching_class", 200), ("academic_year", 100)):
            if field in request.data:
                value = str(request.data.get(field, "")).strip()
                if field != "academic_year" and not value:
                    return _error("课程名称和教学班名称不能为空。", "CLASS_INVALID", status.HTTP_400_BAD_REQUEST)
                if len(value) > limit:
                    return _error("教学班信息长度超过限制。", "CLASS_VALUE_TOO_LONG", status.HTTP_400_BAD_REQUEST)
                update_data[field] = value
        if "is_active" in request.data:
            try:
                update_data["is_active"] = _boolean(request.data.get("is_active"))
            except ValueError:
                return _error("教学班状态参数无效。", "CLASS_STATUS_INVALID", status.HTTP_400_BAD_REQUEST)

        merged = {**current, **update_data}
        duplicate = repository.find_existing(
            str(merged["title"]),
            str(merged["teaching_class"]),
            str(merged.get("academic_year", "")),
            exclude_id=class_id,
        )
        if duplicate:
            return Response(
                {"message": "相同课程、教学班和学年的记录已存在。", "code": "CLASS_EXISTS", "class": duplicate},
                status=status.HTTP_409_CONFLICT,
            )
        repository.update_class(class_id, update_data)
        item = repository.get_class(class_id)
    with LearningRepository() as audit:
        audit.write_audit(session_user(request), "class.update", "classes", class_id, update_data)
    return Response({"class": item})


@api_view(["GET", "POST"])
@permission_classes([IsAdmin])
def class_members(request, class_id: str):
    if request.method == "GET":
        with AdminClassRepository() as repository:
            item = repository.get_class(class_id)
            members = repository.list_members(class_id) if item else None
            teachers = repository.account_options("teacher")
            students = repository.account_options("student")
        if item is None:
            return _error("教学班不存在。", "CLASS_NOT_FOUND", status.HTTP_404_NOT_FOUND)
        return Response({"class": item, "members": members, "options": {"teachers": teachers, "students": students}})

    role = str(request.data.get("role", "")).strip()
    usernames = request.data.get("usernames", [])
    if role not in {"teacher", "student"} or not isinstance(usernames, list):
        return _error("关系参数无效。", "CLASS_MEMBER_INVALID", status.HTTP_400_BAD_REQUEST)
    with AdminClassRepository() as repository:
        if repository.get_class(class_id) is None:
            return _error("教学班不存在。", "CLASS_NOT_FOUND", status.HTTP_404_NOT_FOUND)
        result = repository.replace_members(class_id, role, usernames)
        members = repository.list_members(class_id)
    if result["invalid"]:
        return Response({"message": "部分账号不存在或已停用。", "code": "CLASS_MEMBER_INVALID", "result": result, "members": members}, status=status.HTTP_400_BAD_REQUEST)
    with LearningRepository() as audit:
        audit.write_audit(session_user(request), "class.members.update", "classes", class_id, {"role": role, **result})
    return Response({"result": result, "members": members})
